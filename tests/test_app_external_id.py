"""External document ids through the CSV/XLSX import, results, export and
gold-label round trip (issue #4, first external use case: bbsia-radar).

A corpus imported from a spreadsheet can name one column as the caller's own
stable id. Decifra keeps it as `DocumentRecord.external_id`, shows it next to
every result, and accepts it back as the join key for gold labels, so a
caller never has to match results to its rows by position."""

import csv
import io
import json

import openpyxl
from fastapi.testclient import TestClient
from sqlmodel import Session, select

from text_as_data.app import app, get_engine_dependency, get_provider_dependency
from text_as_data.db import DocumentRecord, HumanLabelRecord, get_engine
from text_as_data.providers import Provider, ProviderResult

VALID_SPEC = {
    "concept": "protest",
    "description": "A collective public event.",
    "categories": [
        {"label": "protest", "definition": "An occupation, march, or strike."},
        {"label": "not_protest", "definition": "Any event that does not meet the criteria above."},
    ],
}


class FakeProvider(Provider):
    def extract(self, messages, schema):
        parsed = schema(category="protest", rationale="because", evidence_span="quote")
        return ProviderResult(parsed=parsed, prompt="fake prompt", raw_response="fake raw response")


def _client():
    engine = get_engine("sqlite://")
    app.dependency_overrides[get_engine_dependency] = lambda: engine
    app.dependency_overrides[get_provider_dependency] = lambda: FakeProvider()
    return TestClient(app), engine


def _upload_csv(client, content: bytes, **form):
    data = {"name": "radar", "text_column": "text", **form}
    return client.post("/corpora/csv", data=data, files={"file": ("corpus.csv", content, "text/csv")})


def _xlsx(rows: list[list]) -> bytes:
    workbook = openpyxl.Workbook()
    for row in rows:
        workbook.active.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _run(client) -> int:
    codebook_id = client.post("/codebooks", json=VALID_SPEC).json()["id"]
    response = client.post("/runs", json={"codebook_id": codebook_id, "corpus_id": "radar", "model": "fake-model"})
    return response.json()["run_id"]


# --- import -----------------------------------------------------------------


def test_csv_import_stores_the_id_column_as_external_id():
    client, engine = _client()
    content = "id,text\ngithub:101,Primeiro documento\nhuggingface:neuralmind/bert,Segundo documento\n".encode()

    response = _upload_csv(client, content, id_column="id")

    assert response.status_code == 200
    assert response.json() == {"corpus_id": "radar", "document_count": 2}
    with Session(engine) as session:
        docs = session.exec(select(DocumentRecord).where(DocumentRecord.corpus_id == "radar")).all()
        assert {d.external_id: d.text for d in docs} == {
            "github:101": "Primeiro documento",
            "huggingface:neuralmind/bert": "Segundo documento",
        }


def test_csv_import_without_id_column_keeps_external_id_empty():
    client, engine = _client()

    response = _upload_csv(client, b"id,text\na,one\nb,two\n")

    assert response.status_code == 200
    with Session(engine) as session:
        assert all(d.external_id is None for d in session.exec(select(DocumentRecord)).all())


def test_csv_import_rejects_unknown_id_column_and_creates_nothing():
    client, engine = _client()

    response = _upload_csv(client, b"id,text\na,one\n", id_column="solution_id")

    assert response.status_code == 422
    assert "solution_id" in response.json()["detail"]
    with Session(engine) as session:
        assert session.exec(select(DocumentRecord)).all() == []


def test_csv_import_rejects_duplicate_external_ids_and_names_them():
    client, engine = _client()

    response = _upload_csv(client, b"id,text\na,one\nb,two\na,three\n", id_column="id")

    assert response.status_code == 422
    assert "'a'" in response.json()["detail"]
    with Session(engine) as session:
        assert session.exec(select(DocumentRecord)).all() == []


def test_csv_import_rejects_blank_external_id_on_a_kept_row():
    client, _ = _client()

    response = _upload_csv(client, b"id,text\na,one\n  ,two\n", id_column="id")

    assert response.status_code == 422
    assert "blank" in response.json()["detail"]


def test_rows_skipped_for_empty_text_do_not_count_as_duplicate_or_blank_ids():
    client, engine = _client()

    response = _upload_csv(client, b"id,text\na,one\na,\n,\n", id_column="id")

    assert response.status_code == 200
    assert response.json()["document_count"] == 1
    with Session(engine) as session:
        assert [d.external_id for d in session.exec(select(DocumentRecord)).all()] == ["a"]


def test_external_id_is_trimmed():
    client, engine = _client()

    _upload_csv(client, b"id,text\n  a  ,one\n", id_column="id")

    with Session(engine) as session:
        assert session.exec(select(DocumentRecord)).one().external_id == "a"


def test_xlsx_import_turns_integral_numeric_ids_into_plain_strings():
    client, engine = _client()
    content = _xlsx([["id", "text"], [101, "one"], [102.0, "two"], ["x-3", "three"]])

    response = client.post(
        "/corpora/xlsx",
        data={"name": "radar", "text_column": "text", "id_column": "id"},
        files={"file": ("corpus.xlsx", content, "application/octet-stream")},
    )

    assert response.status_code == 200
    with Session(engine) as session:
        assert sorted(d.external_id for d in session.exec(select(DocumentRecord)).all()) == ["101", "102", "x-3"]


# --- results and export -------------------------------------------------------


def test_results_and_every_export_format_carry_the_external_id():
    client, _ = _client()
    _upload_csv(client, b"id,text\ngithub:101,one\ngithub:102,two\n", id_column="id")
    run_id = _run(client)

    results = client.get(f"/runs/{run_id}/results").json()
    assert sorted(r["document_external_id"] for r in results) == ["github:101", "github:102"]

    exported_json = json.loads(client.get(f"/runs/{run_id}/export?format=json").content)
    assert sorted(r["document_external_id"] for r in exported_json) == ["github:101", "github:102"]

    exported_csv = client.get(f"/runs/{run_id}/export?format=csv").content.decode("utf-8-sig")
    rows = list(csv.DictReader(io.StringIO(exported_csv)))
    assert sorted(r["document_external_id"] for r in rows) == ["github:101", "github:102"]

    workbook = openpyxl.load_workbook(io.BytesIO(client.get(f"/runs/{run_id}/export?format=xlsx").content))
    header = [c.value for c in next(workbook.active.iter_rows(max_row=1))]
    assert "document_external_id" in header


def test_results_show_empty_external_id_for_corpora_imported_without_one():
    client, _ = _client()
    client.post("/corpora/paste", json={"name": "radar", "text": "pasted"})
    run_id = _run(client)

    results = client.get(f"/runs/{run_id}/results").json()

    assert results[0]["document_external_id"] is None


# --- gold labels by external id -----------------------------------------------


def test_gold_labels_can_be_matched_by_external_id():
    client, engine = _client()
    _upload_csv(client, b"id,text\ngithub:101,one\ngithub:102,two\n", id_column="id")
    run_id = _run(client)
    gold = b"document_external_id,gold_category\ngithub:101,not_protest\ngithub:102,\n"

    response = client.post(f"/runs/{run_id}/gold-labels", files={"file": ("gold.csv", gold, "text/csv")})

    assert response.status_code == 200
    assert response.json() == {"imported": 1, "skipped_blank": 1}
    with Session(engine) as session:
        label = session.exec(select(HumanLabelRecord)).one()
        document = session.get(DocumentRecord, label.document_id)
        assert (document.external_id, label.category) == ("github:101", "not_protest")


def test_gold_labels_reject_external_ids_outside_the_run_corpus():
    client, _ = _client()
    _upload_csv(client, b"id,text\ngithub:101,one\n", id_column="id")
    run_id = _run(client)
    gold = b"document_external_id,gold_category\ngithub:999,protest\n"

    response = client.post(f"/runs/{run_id}/gold-labels", files={"file": ("gold.csv", gold, "text/csv")})

    assert response.status_code == 422
    assert "github:999" in response.json()["detail"]


def test_gold_labels_prefer_document_id_when_both_columns_are_present():
    """A file exported from Decifra has both columns; the internal id stays
    authoritative so older gold files keep working unchanged."""
    client, engine = _client()
    _upload_csv(client, b"id,text\ngithub:101,one\n", id_column="id")
    run_id = _run(client)
    document_id = client.get(f"/runs/{run_id}/results").json()[0]["document_id"]
    gold = f"document_id,document_external_id,gold_category\n{document_id},github:101,protest\n".encode()

    response = client.post(f"/runs/{run_id}/gold-labels", files={"file": ("gold.csv", gold, "text/csv")})

    assert response.status_code == 200
    with Session(engine) as session:
        assert session.exec(select(HumanLabelRecord)).one().document_id == document_id


def test_gold_labels_still_require_an_id_column():
    client, _ = _client()
    _upload_csv(client, b"id,text\ngithub:101,one\n", id_column="id")
    run_id = _run(client)

    response = client.post(
        f"/runs/{run_id}/gold-labels", files={"file": ("gold.csv", b"gold_category\nprotest\n", "text/csv")}
    )

    assert response.status_code == 422
    assert "document_external_id" in response.json()["detail"]


def test_gold_labels_round_trip_ids_that_the_export_defuses_as_formulas():
    """The CSV export prefixes `-1` with an apostrophe (CSV-injection guard).
    Re-uploading that exported value as a gold key must still match `-1`."""
    client, engine = _client()
    _upload_csv(client, b"id,text\n-1,one\n@team,two\n", id_column="id")
    run_id = _run(client)
    exported = client.get(f"/runs/{run_id}/export?format=csv").content.decode("utf-8-sig")
    rows = list(csv.DictReader(io.StringIO(exported)))
    assert sorted(r["document_external_id"] for r in rows) == ["'-1", "'@team"]
    gold = "document_external_id,gold_category\n" + "".join(
        f"{r['document_external_id']},protest\n" for r in rows
    )

    response = client.post(
        f"/runs/{run_id}/gold-labels", files={"file": ("gold.csv", gold.encode(), "text/csv")}
    )

    assert response.status_code == 200
    assert response.json()["imported"] == 2
    with Session(engine) as session:
        labeled = {session.get(DocumentRecord, label.document_id).external_id for label in session.exec(select(HumanLabelRecord))}
        assert labeled == {"-1", "@team"}


def test_blank_id_errors_count_data_rows_not_spreadsheet_lines():
    """The XLSX parser drops fully empty rows, so errors name the position
    among data rows instead of a spreadsheet line that would be off."""
    client, _ = _client()
    content = _xlsx([["id", "text"], ["a", "one"], [None, None], [None, "two"]])

    response = client.post(
        "/corpora/xlsx",
        data={"name": "radar", "text_column": "text", "id_column": "id"},
        files={"file": ("corpus.xlsx", content, "application/octet-stream")},
    )

    assert response.status_code == 422
    assert "data row(s) [2]" in response.json()["detail"]
